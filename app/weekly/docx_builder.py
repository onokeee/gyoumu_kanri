"""週報のWord(.docx)ファイルを作る(python-docx)。

文書の構成(1回の作成で1ファイル):
  タイトル「週報」／対象期間／「(アプリ名)が自動作成（作成日時）」
  チーム全体 : 見出し＋文章＋メンバー別の集計表(コードで作成)＋チーム合計の1行
  個人       : 改ページして1人1ページ(氏名の見出し＋文章)

文章(AI または ルールベース)は1行ずつ次のように変換する:
  ■ または 【 で始まる行 → 見出し2
  ・ - * で始まる行      → 箇条書き(List Bullet)
  空行                   → 飛ばす
  Markdown の記号(行頭の #・**)は取り除く
Word(XML)に入れられない制御文字(貼り付けた端末出力の ESC など)は、すべての文字列から取り除く
(1文字でも残っていると文書全体が作成できないため)。
"""
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.weekly.rules import period_label

FONT_NAME = "Yu Gothic"
FONT_SIZE = Pt(10.5)

TABLE_HEADERS = ["氏名", "完了", "進行中", "期限超過", "進捗記載数", "月間負荷h", "余力h"]

_MD_HEADING = re.compile(r"^#{1,6}\s*")
_MD_RULE = re.compile(r"^[-=_*]{3,}$")
_BULLET_MARKS = ("・", "-", "*", "•")
_HEADING_MARKS = ("■", "【")
# XML 1.0 で使えない文字(タブ・改行以外の制御文字、サロゲート、U+FFFE/U+FFFF)
_XML_INVALID = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _xml_text(value):
    """Word に入れられない制御文字などを取り除いた文字列にする。"""
    return _XML_INVALID.sub("", str(value or ""))


def _set_style_font(style, size=None):
    """スタイルのフォントを游ゴシックにする(日本語用の東アジアフォントも指定)。"""
    style.font.name = FONT_NAME
    if size is not None:
        style.font.size = size
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), FONT_NAME)
    # テーマのフォント指定が残っていると明示したフォントより優先されるため外す
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)


def _new_document(title):
    doc = Document()
    _set_style_font(doc.styles["Normal"], FONT_SIZE)
    for name in ("Title", "Heading 1", "Heading 2", "List Bullet"):
        try:
            _set_style_font(doc.styles[name])
        except KeyError:
            pass

    # A4・余白20mm
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(20))

    doc.core_properties.title = title
    return doc


def _add_note(doc, text):
    """小さめの灰色の注記(「AI未整形」など)。"""
    run = doc.add_paragraph().add_run("※" + _xml_text(text))
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6C, 0x75, 0x7D)


def add_text(doc, text):
    """文章を1行ずつ見出し・箇条書き・本文に変換して追加する。"""
    for raw in (text or "").splitlines():
        line = _xml_text(raw).strip().replace("**", "")
        is_md_heading = bool(_MD_HEADING.match(line))
        line = _MD_HEADING.sub("", line).strip()
        if not line or _MD_RULE.match(line):
            continue
        if is_md_heading or line.startswith(_HEADING_MARKS):
            doc.add_paragraph(line, style="Heading 2")
        elif line.startswith(_BULLET_MARKS):
            body = line[1:].strip()
            if body:
                doc.add_paragraph(body, style="List Bullet")
        else:
            doc.add_paragraph(line)


def _add_member_table(doc, rows):
    """メンバー別の集計表(コードで集計した値をそのまま載せる)。"""
    table = doc.add_table(rows=1, cols=len(TABLE_HEADERS))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, TABLE_HEADERS):
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in rows:
        values = [
            r["name"], str(r["completed"]), str(r["doing"]), str(r["overdue"]),
            str(r["comments"]), "{:.1f}".format(r["total_h"]), "{:.1f}".format(r["spare"]),
        ]
        for i, (cell, value) in enumerate(zip(table.add_row().cells, values)):
            cell.text = _xml_text(value)
            if i > 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return table


def _team_total_line(team):
    return (
        "チーム合計（複数担当のタスクは1件）: 完了 {}件・着手 {}件・新規 {}件・"
        "進行中 {}件・保留 {}件・期限超過 {}件 ／ 成果（期間内に完了・年換算）: "
        "金額 {:,.0f} ￥/年・時間 {:,.1f} ｈ/年"
    ).format(
        team["completed"], team["started"], team["created"], team["doing"],
        team["hold"], team["overdue"], team["money_act"], team["hour_act"],
    )


def build_docx(material, written, created_at, app_name):
    """週報の .docx を作ってバイト列で返す。

    material : collector.collect() の戻り値(期間・チーム集計)
    written  : writer.write_report() の戻り値(チーム全体・各人の文章)
    """
    period = period_label(material["start"], material["end"])
    doc = _new_document("週報 {}".format(period))

    doc.add_heading("週報", level=0)
    doc.add_paragraph("対象期間: {}（対象 {}名）".format(period, len(written["persons"])))
    info = doc.add_paragraph().add_run(
        "{}が自動作成（{}）".format(_xml_text(app_name), created_at.strftime("%Y/%m/%d %H:%M")))
    info.font.size = Pt(9)
    info.font.color.rgb = RGBColor(0x6C, 0x75, 0x7D)

    # ---- チーム全体 ----
    doc.add_heading("チーム全体", level=1)
    team_text = written["team"]
    if team_text["note"]:
        _add_note(doc, team_text["note"])
    add_text(doc, team_text["text"])

    doc.add_paragraph("■メンバー別の集計（システム集計）", style="Heading 2")
    _add_member_table(doc, material["team"]["rows"])
    total = doc.add_paragraph().add_run(_team_total_line(material["team"]))
    total.font.size = Pt(9)

    # ---- 個人(1人1ページ) ----
    for person in written["persons"]:
        doc.add_page_break()
        doc.add_heading(_xml_text(person["name"]), level=1)
        if person["note"]:
            _add_note(doc, person["note"])
        add_text(doc, person["text"])

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()
