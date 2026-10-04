"""全データの Excel(.xlsx) 出力。マネージャーのみ。

各メニューのデータ(現在データ＋履歴データ)を openpyxl で xlsx 化し、
添付ファイルとしてダウンロードさせる。読み取り専用(DBは変更しない)。
"""
from datetime import datetime, date
from io import BytesIO

from flask import Blueprint, Response, abort
from flask_login import login_required, current_user
from openpyxl import Workbook
from openpyxl.styles import Font

from app.models.user import User
from app.models.task import Task, TaskComment, TaskStatusChange
from app.models.routine import RoutineWork
from app.models.skill import Skill, SkillRating
from app.models.operation import Operation, OperationSkill
from app.models.leave import LeaveRequest
from app.models.department import Department
from app.models.skilltest import (
    SkillTestAnswer,
    SkillTestAttempt,
    SkillTestQuestion,
    choice_letter,
)

export_bp = Blueprint("export", __name__, url_prefix="/export")


@export_bp.before_request
@login_required
def _managers_only():
    """データ出力はマネージャーのみ。"""
    if not current_user.is_manager:
        abort(403)


def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "はい" if v else "いいえ"
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return v


def _xlsx_response(sheets, filename):
    """sheets = [(シート名, ヘッダ list, 行 list of list)] を1つの xlsx にして返す。"""
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, rows in sheets:
        ws = wb.create_sheet(title=name[:31])
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_fmt(c) for c in row])
            # openpyxl は「=」で始まる文字列を数式として保存するため、文字列として書き出す
            # (タスク名・コメント・AIが作った問題文や選択肢などが数式として計算されないように)
            for cell in ws[ws.max_row]:
                if cell.data_type == "f" and isinstance(cell.value, str):
                    cell.data_type = "s"
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = \
                max(10, min(45, len(str(h)) * 2 + 6))
    if not wb.sheetnames:
        wb.create_sheet(title="data")
    bio = BytesIO()
    wb.save(bio)
    return Response(
        bio.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# --------------------------------------------------------------------------- #
# 各メニューのシート定義(現在データ＋履歴データ)
# --------------------------------------------------------------------------- #
def _tasks_sheets():
    tasks = Task.query.order_by(Task.id).all()
    t_headers = [
        "ID", "タイトル", "ステータス", "優先度", "開始日", "期限", "規模",
        "担当者", "登録者", "成果定量-見込み", "単位", "成果定量-実績", "単位",
        "成果-補足", "成果定性-見込み", "成果定性-実績", "内容・詳細",
        "登録日時", "更新日時",
    ]
    t_rows = [[
        t.id, t.title, t.status, t.priority, t.start_date, t.due_date,
        t.scale_label or "", t.assignee_names,
        t.creator.display_name if t.creator else "",
        t.outcome_quant_estimate, t.outcome_quant_estimate_unit,
        t.outcome_quant_actual, t.outcome_quant_actual_unit,
        t.outcome_quant_note, t.outcome_qual_estimate, t.outcome_qual_actual,
        t.description, t.created_at, t.updated_at,
    ] for t in tasks]

    changes = TaskStatusChange.query.order_by(
        TaskStatusChange.task_id, TaskStatusChange.changed_at
    ).all()
    c_rows = [[
        ch.task_id, ch.task.title if ch.task else "", ch.status, ch.changed_at,
    ] for ch in changes]

    comments = TaskComment.query.order_by(
        TaskComment.task_id, TaskComment.created_at
    ).all()
    cm_rows = [[
        cm.task_id, cm.task.title if cm.task else "",
        cm.user.display_name if cm.user else "", cm.body, cm.created_at,
    ] for cm in comments]

    return [
        ("タスク(現在)", t_headers, t_rows),
        ("ステータス履歴", ["タスクID", "タスク", "ステータス", "変更日時"], c_rows),
        ("進捗状況(履歴)", ["タスクID", "タスク", "記載者", "進捗内容", "記載日時"], cm_rows),
    ]


def _routine_sheets():
    rows = [[
        r.id, r.name, r.assignee.display_name if r.assignee else "",
        r.purpose, r.frequency_count, r.frequency_unit, r.minutes_per,
        r.monthly_minutes, r.content, r.manual_status,
        r.creator.display_name if r.creator else "", r.created_at, r.updated_at,
    ] for r in RoutineWork.query.order_by(RoutineWork.assignee_id, RoutineWork.id).all()]
    headers = [
        "ID", "業務名", "担当者", "目的", "回数", "頻度単位", "1回所要(分)",
        "月間所要(分)", "業務内容", "手順書", "登録者", "登録日時", "更新日時",
    ]
    return [("定型・定期業務", headers, rows)]


def _skills_sheets():
    skills = Skill.query.order_by(Skill.skill_type, Skill.sort_order, Skill.name).all()
    s_rows = [[
        s.id, s.name, s.type_label, s.category, s.description, s.is_active, s.sort_order,
    ] for s in skills]

    ratings = SkillRating.query.order_by(SkillRating.user_id, SkillRating.skill_id).all()
    r_rows = [[
        r.user.display_name if r.user else "", r.skill.name if r.skill else "",
        r.level, r.note, r.rater.display_name if r.rater else "", r.rated_at,
    ] for r in ratings]

    ops = Operation.query.order_by(Operation.sort_order, Operation.name).all()
    o_rows = [[o.id, o.name, o.description, o.is_active, o.sort_order] for o in ops]

    reqs = OperationSkill.query.all()
    req_rows = [[
        r.operation.name if r.operation else "",
        r.skill.name if r.skill else "", r.level,
    ] for r in reqs]

    return [
        ("スキル項目", ["ID", "スキル名", "区分", "カテゴリ", "説明", "有効", "並び順"], s_rows),
        ("到達度", ["メンバー", "スキル", "到達度", "メモ", "評価者", "評価日時"], r_rows),
        ("業務", ["ID", "業務名", "説明", "有効", "並び順"], o_rows),
        ("業務別-必要スキル", ["業務", "必要スキル", "必要レベル"], req_rows),
    ]


def _leaves_sheets():
    rows = [[
        lv.id, lv.user.display_name if lv.user else "",
        lv.leave_date, lv.leave_type, lv.day_count, lv.created_at,
    ] for lv in LeaveRequest.query.order_by(LeaveRequest.leave_date).all()]
    headers = ["ID", "メンバー", "取得日", "種別", "換算日数", "登録日時"]
    return [("年休", headers, rows)]


def _skilltest_sheets():
    """スキルテストの受験履歴・回答(全問)・問題プール。"""
    attempts = SkillTestAttempt.query.order_by(SkillTestAttempt.started_at, SkillTestAttempt.id).all()
    a_headers = [
        "受験ID", "メンバー", "スキル", "状態", "開始日時", "終了日時", "全体の期限",
        "問題数", "正解数", "正答率(%)", "レベル別(正解/問題・判定)", "結果レベル",
        "受験前の到達度", "受験後の到達度", "自動登録", "離脱回数", "再出題数", "受験時の設定",
    ]
    a_rows = []
    for a in attempts:
        per_level = "、".join(
            "Lv{} {}/{}{}".format(
                r.get("level"), r.get("correct"), r.get("total"), "合格" if r.get("passed") else "")
            for r in a.level_result_list
        )
        a_rows.append([
            a.id, a.user.display_name if a.user else "", a.skill.name if a.skill else "",
            a.status_label, a.started_at, a.finished_at, a.deadline_at,
            a.total, None if a.is_in_progress else a.correct,
            None if a.is_in_progress else a.rate, per_level, a.result_level,
            a.prev_level, a.new_level, None if a.is_in_progress else a.applied,
            a.blur_count, a.reused_count, a.settings_snapshot or "",
        ])

    answers = (
        SkillTestAnswer.query.join(SkillTestAttempt)
        .order_by(SkillTestAttempt.started_at, SkillTestAnswer.attempt_id, SkillTestAnswer.seq)
        .all()
    )
    w_headers = [
        "受験ID", "メンバー", "スキル", "出題順", "レベル", "問題ID", "問題文",
        "選択肢A", "選択肢B", "選択肢C", "選択肢D", "正解", "回答", "結果",
        "制限時間(秒)", "表示日時", "確定日時", "所要(秒)", "離脱回数",
    ]
    w_rows = []
    for w in answers:
        choices = (w.choice_list + ["", "", "", ""])[:4]
        attempt = w.attempt
        w_rows.append([
            w.attempt_id,
            attempt.user.display_name if attempt and attempt.user else "",
            attempt.skill.name if attempt and attempt.skill else "",
            w.seq, w.level, w.question_id, w.question,
            choices[0], choices[1], choices[2], choices[3],
            choice_letter(w.correct_index), choice_letter(w.selected_index), w.result_label,
            w.time_limit_sec, w.served_at, w.answered_at, w.elapsed_sec, w.blur_count,
        ])

    questions = SkillTestQuestion.query.order_by(
        SkillTestQuestion.skill_id, SkillTestQuestion.level, SkillTestQuestion.id).all()
    q_headers = [
        "問題ID", "スキル", "レベル", "問題文", "選択肢A", "選択肢B", "選択肢C", "選択肢D",
        "正解", "解説", "作成元", "モデル", "有効", "作成日時",
    ]
    q_rows = []
    for q in questions:
        choices = (q.choice_list + ["", "", "", ""])[:4]
        q_rows.append([
            q.id, q.skill.name if q.skill else "", q.level, q.question,
            choices[0], choices[1], choices[2], choices[3],
            q.answer_letter, q.explanation, q.source_label, q.model, q.is_active, q.created_at,
        ])

    return [
        ("スキルテスト受験履歴", a_headers, a_rows),
        ("スキルテスト回答", w_headers, w_rows),
        ("スキルテスト問題", q_headers, q_rows),
    ]


def _teams_sheets():
    depts = Department.query.order_by(Department.sort_order, Department.name).all()
    d_rows = [[
        d.id, d.name, d.is_active, d.sort_order,
        "、".join(u.display_name for u in d.users),
    ] for d in depts]

    users = User.query.order_by(User.display_name).all()
    u_rows = [[
        u.username, u.display_name, u.role_label, u.is_active, u.department_names,
    ] for u in users]

    return [
        ("チーム", ["ID", "チーム名", "有効", "並び順", "所属メンバー"], d_rows),
        ("メンバー", ["ログインID", "氏名", "役割", "有効", "所属チーム"], u_rows),
    ]


# --------------------------------------------------------------------------- #
# ダウンロード用ルート
# --------------------------------------------------------------------------- #
@export_bp.route("/tasks.xlsx")
def tasks_xlsx():
    return _xlsx_response(_tasks_sheets(), "tasks.xlsx")


@export_bp.route("/routine.xlsx")
def routine_xlsx():
    return _xlsx_response(_routine_sheets(), "routine.xlsx")


@export_bp.route("/skills.xlsx")
def skills_xlsx():
    return _xlsx_response(_skills_sheets(), "skills.xlsx")


@export_bp.route("/leaves.xlsx")
def leaves_xlsx():
    return _xlsx_response(_leaves_sheets(), "leaves.xlsx")


@export_bp.route("/skilltest.xlsx")
def skilltest_xlsx():
    return _xlsx_response(_skilltest_sheets(), "skilltest.xlsx")


@export_bp.route("/teams.xlsx")
def teams_xlsx():
    return _xlsx_response(_teams_sheets(), "teams.xlsx")


@export_bp.route("/all.xlsx")
def all_xlsx():
    sheets = (
        _tasks_sheets() + _routine_sheets() + _skills_sheets()
        + _skilltest_sheets() + _leaves_sheets() + _teams_sheets()
    )
    return _xlsx_response(sheets, "all_data.xlsx")
